#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""validate_minutes.py — 结构化总结结构校验与源文核对（meeting-notes-expert 配套）

按会议版章节名校验，仅适用会议场景的会议纪要；讲座、访谈、播客等通用总结不强制。
做四件事：
  1) 结构：必需节齐备（会议结论/关键决策/待办事项/会议内容），按 minutes-schema.json 校验；
  2) 版式：待办表六列、优先级枚举合法、基本信息区不得是表格；
  3) 文风：正文不得出现「待确认/存疑/转写亦作」等禁用标注，决策不得写成「讨论了…」；
  4) 溯源（给了 --source 才做）：成稿中的数字必须出现在源材料；待办责任人必须出现在源材料。

用法：
    python validate_minutes.py 纪要.md [--source 会议原文.txt] [--strict] [--json 报告.json]
    python validate_minutes.py 纪要.docx [--source ...]
    python validate_minutes.py --selftest

退出码：0 通过；1 有 FAIL（--strict 时 WARN 也计 FAIL）；2 用法/读取错误。
"""
import argparse
import json
import re
import sys
from pathlib import Path

TOOL_DIR = Path(__file__).resolve().parent
SCHEMA_PATH = TOOL_DIR.parent / "references" / "minutes-schema.json"

REQUIRED = ["会议结论", "关键决策", "待办事项", "会议内容"]
OPTIONAL = ["待定与开放问题", "上期行动项回顾"]
ALL_SECTIONS = REQUIRED + OPTIONAL
FORBIDDEN = ["待确认", "存疑", "转写亦作", "不确定标注"]
ACTION_COLUMNS = ["序号", "待办事项", "责任人/责任部门", "完成时间", "优先级", "备注"]
PRIORITIES = {"高", "中", "低", "待定"}
BASIC_FIELDS_OK = {"会议主题", "会议时间", "会议地点", "参会人员", "记录人",
                   "主持人", "会议方式", "会议编号", "会议形式"}
HEADING_RE = re.compile(r"^##(?!#)\s+(.+)$")
HEADING_STYLE_RE = re.compile(r"^(?:Heading|标题)\s*(\d+)$")
LIST_RE = re.compile(r"^\s*(?:[-*+]|\d+[.、)])\s+")
NUM_RE = re.compile(r"\d+(?:[.,]\d+)*\s*[%％]?")
OWNER_PLACEHOLDERS = {"待定", "未指定", "全体", "大家", "—", "-"}


# ---------- 解析 ----------
def parse_md(text):
    blocks, table = [], None
    for raw in text.splitlines():
        line = raw.rstrip()
        if line.lstrip().startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if cells and all(set(c) <= set("-: ") for c in cells):
                continue
            if table is None:
                table = []
                blocks.append(("t", table))
            table.append(cells)
            continue
        table = None
        blocks.append(("p", line))
    return blocks


def parse_docx(path):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    blocks = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, doc)
            text = para.text
            # article-format 出稿后标题是样式不是字面 `##`，这里按样式还原成
            # Markdown 标题，让 split_sections 的同一套规则同时吃 md 与 docx。
            style = (para.style.name or "").strip() if para.style is not None else ""
            m = HEADING_STYLE_RE.match(style)
            if m:
                text = "#" * (int(m.group(1)) + 1) + " " + text
            elif style == "Title":
                text = "# " + text
            blocks.append(("p", text))
        elif child.tag == qn("w:tbl"):
            rows = [[c.text.strip() for c in r.cells] for r in Table(child, doc).rows]
            blocks.append(("t", rows))
    return blocks


def split_sections(blocks):
    basic, sections, basic_table = {}, {}, False
    cur = None
    for kind, payload in blocks:
        if kind == "p":
            text = payload.strip()
            m = HEADING_RE.match(text)
            if m:
                name = m.group(1).strip()
                cur = name if name in ALL_SECTIONS else None
                if cur:
                    sections.setdefault(cur, [])
                continue
        if cur is None:
            if kind == "t":
                basic_table = True
            elif not payload.lstrip().startswith("#"):
                m = re.match(r"^([^：:]{2,12})[：:]\s*(.+)$", payload.strip())
                if m:
                    basic[m.group(1).strip()] = m.group(2).strip()
        else:
            sections[cur].append((kind, payload))
    return basic, sections, basic_table


def section_text(entries):
    out = []
    for kind, payload in entries:
        if kind == "p":
            out.append(payload)
        else:
            out.extend(" | ".join(r) for r in payload)
    return "\n".join(out).strip()


def collect_action_items(entries):
    """取待办节的第一张表。列数不对也返回，让列名断言报出来，而不是静默当作没有表。"""
    keys = ["id", "task", "owner", "due", "priority", "note"]
    for kind, payload in entries:
        if kind != "t" or not payload:
            continue
        header = [c.strip() for c in payload[0]]
        items = []
        for row in payload[1:]:
            cells = [c.strip() for c in row]
            items.append({k: (cells[i] if i < len(cells) else "") for i, k in enumerate(keys)})
        return items, header
    return [], []


def collect_decisions(entries):
    lines = []
    for kind, payload in entries:
        if kind != "p":
            continue
        text = payload.strip()
        if text:
            lines.append(LIST_RE.sub("", text).strip())
    return [x for x in lines if x]


# ---------- 校验 ----------
def validate(path, source_text=None, strict=False):
    results = []

    def ok(name, detail=""):
        results.append(("PASS", name, detail))

    def bad(name, detail=""):
        results.append(("FAIL", name, detail))

    def warn(name, detail=""):
        results.append(("WARN", name, detail))

    try:
        text = path.read_text(encoding="utf-8")
        blocks = parse_md(text)
    except (UnicodeDecodeError, OSError):
        try:
            blocks = parse_docx(path)
        except Exception as e:
            print(f"读取失败：{e}")
            return 2, []

    basic, sections, basic_table = split_sections(blocks)
    if basic_table:
        bad("基本信息表形态", "基本信息区使用了表格，规范要求逐行「字段名：内容」")
    if not basic:
        warn("基本信息", "未解析到字段行（材料可能未提供基本信息）")
    else:
        unknown = [k for k in basic if k not in BASIC_FIELDS_OK]
        (warn if unknown else ok)("基本信息字段", "非常规字段：" + "、".join(unknown) if unknown else "")

    missing = [s for s in REQUIRED if not section_text(sections.get(s, []))]
    if missing:
        bad("必需节齐备", "缺：" + "、".join(missing))
    else:
        ok("必需节齐备", "结论/决策/待办/内容 均在")

    body = "\n".join(section_text(v) for v in sections.values())
    hits = [w for w in FORBIDDEN if w in body]
    if hits:
        bad("禁用标注", "正文出现：" + "、".join(hits))
    else:
        ok("禁用标注", "无待确认/存疑类标注")

    decisions = collect_decisions(sections.get("关键决策", []))
    if not decisions:
        bad("关键决策", "节内无条目（未形成决策时应写「本次会议未形成决策」）")
    else:
        bad_lines = [d for d in decisions if re.match(r"^(讨论|研究|考虑)了", d)]
        if bad_lines:
            bad("决策句式", "以下条目是讨论句式而非结论：" + "；".join(bad_lines[:3]))
        else:
            ok("决策句式", f"{len(decisions)} 条均为肯定表述")

    items, header = collect_action_items(sections.get("待办事项", []))
    if header:
        if header == ACTION_COLUMNS:
            ok("待办表列名", "六列且顺序正确")
        else:
            bad("待办表列名", "应为 " + "/".join(ACTION_COLUMNS) + "，实为 " + "/".join(header))
    elif "未形成行动事项" in section_text(sections.get("待办事项", [])):
        ok("待办表", "声明未形成行动事项")
    else:
        bad("待办表", "待办节既无六列表格，也无「未形成行动事项」声明")

    if items:
        for i, it in enumerate(items, 1):
            if not it["task"]:
                bad(f"待办第{i}行", "事项为空")
            if not it["owner"]:
                bad(f"待办第{i}行", "责任人列缺失（无责任人应填「待定」）")
            if not it["due"]:
                bad(f"待办第{i}行", "完成时间列缺失（未定时应填「待定」）")
            if it["priority"] not in PRIORITIES:
                bad(f"待办第{i}行", f"优先级「{it['priority']}」不在 高/中/低/待定 内")
        if not any(r[0] == "FAIL" and r[1].startswith("待办第") for r in results):
            ok("待办字段", f"{len(items)} 行责任人与时间均非空")

    if not missing:
        try:
            import jsonschema
            schema = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))
            payload = {
                "basic_info": basic,
                "sections": {k: section_text(v) for k, v in sections.items()},
                "action_items": items if (items and header == ACTION_COLUMNS) else [],
                "decisions": decisions,
            }
            jsonschema.validate(payload, schema)
            ok("结构 schema", "符合 minutes-schema.json")
        except ImportError:
            warn("结构 schema", "jsonschema 未安装，跳过")
        except FileNotFoundError:
            warn("结构 schema", "minutes-schema.json 缺失，跳过")
        except Exception as e:
            warn("结构 schema", str(e).splitlines()[0][:120])

    # ---- 溯源核对（可选）----
    if source_text:
        src = source_text
        src_nums = {re.sub(r"[,\s]", "", x).rstrip("%％") for x in NUM_RE.findall(src)}
        out_nums = {re.sub(r"[,\s]", "", x).rstrip("%％") for x in NUM_RE.findall(body)}
        invented = sorted(n for n in out_nums if n not in src_nums and len(n) >= 2)
        if invented:
            warn("数字溯源", "源材料未出现：" + "、".join(invented[:8]))
        else:
            ok("数字溯源", f"{len(out_nums)} 个数字均可在源材料中找到")

        owners = [it["owner"] for it in items
                  if it["owner"] and it["owner"] not in OWNER_PLACEHOLDERS]
        unknown_owners = [o for o in owners if o not in src]
        if unknown_owners:
            warn("责任人溯源", "源材料未出现：" + "、".join(sorted(set(unknown_owners))))
        elif owners:
            ok("责任人溯源", f"{len(set(owners))} 位责任人均在源材料中出现")
        else:
            ok("责任人溯源", "无可核对的责任人（均为待定/未指定）")

    fails = sum(1 for r in results if r[0] == "FAIL")
    warns = sum(1 for r in results if r[0] == "WARN")
    if strict:
        fails += warns
    return (1 if fails else 0), results


def summarize(results):
    """三态判定：complete 全过 / degraded 无 FAIL 但有 WARN / failed 有 FAIL。

    借 silverstein/minutes 的 status 枚举思路——把「结构可用但存在非致命问题」
    从二元的 PASS/FAIL 里分出来，避免有 WARN 的稿子被当成完全健康或彻底报废。
    """
    fails = [r for r in results if r[0] == "FAIL"]
    warns = [r for r in results if r[0] == "WARN"]
    if fails:
        status = "failed"
    elif warns:
        status = "degraded"
    else:
        status = "complete"
    warnings = [{"step": name, "reason": detail or "未给原因", "message": detail}
                for _, name, detail in warns]
    return status, warnings


SELFTEST_OK = """# 会议纪要：测试会议

## 会议基本信息

会议主题：测试
会议时间：2026-10-19
参会人员：张三、李四

## 会议结论

本次会议决定采用方案 A。

## 关键决策

决定采用方案 A 上线。

## 待办事项

| 序号 | 待办事项 | 责任人/责任部门 | 完成时间 | 优先级 | 备注 |
|---|---|---|---|---|---|
| 1 | 完成上线检查 | 张三 | 10 月 25 日 | 高 | 优先级未明示 |

## 待定与开放问题

预算口径待定，由李四跟进。

## 会议内容

### 上线方案

讨论后确定采用方案 A。
"""

SELFTEST_BAD = """# 会议纪要：坏例

## 会议结论

待确认。

## 关键决策

讨论了方案 A。

## 待办事项

| 序号 | 待办事项 | 责任人 | 完成时间 |
|---|---|---|---|
| 1 | 处理一下 |  |  |

## 会议内容

无。
"""


def selftest():
    import tempfile
    ok = True
    with tempfile.TemporaryDirectory() as td:
        good = Path(td) / "good.md"
        good.write_text(SELFTEST_OK, encoding="utf-8")
        rc, res = validate(good)
        print(f"[{'PASS' if rc == 0 else 'FAIL'}] 正例应通过（exit={rc}）")
        ok &= rc == 0
        for level, name, detail in res:
            if level != "PASS":
                print(f"  {level} {name} {detail}")
        badp = Path(td) / "bad.md"
        badp.write_text(SELFTEST_BAD, encoding="utf-8")
        rc2, res2 = validate(badp)
        caught = [r for r in res2 if r[0] == "FAIL"]
        print(f"[{'PASS' if rc2 == 1 and len(caught) >= 4 else 'FAIL'}] 负例应被抓（exit={rc2}，{len(caught)} 项 FAIL）")
        ok &= rc2 == 1 and len(caught) >= 4
        # 数字溯源负例：成稿出现源材料没有的数字
        src = "本次会议由张三汇报，产量 120 吨。"
        rc3, res3 = validate(good, source_text=src)
        warn_hit = any(r[0] == "WARN" and r[1] == "数字溯源" for r in res3)
        print(f"[{'PASS' if warn_hit else 'FAIL'}] 编造数字应告警")
        ok &= warn_hit
        # 三态判定：干净稿=complete，有 WARN=degraded，有 FAIL=failed
        st_clean, _ = summarize(validate(good)[1])
        st_warn, wl = summarize(res3)
        st_fail, _ = summarize(validate(badp)[1])
        tri_ok = st_clean == "complete" and st_warn == "degraded" and st_fail == "failed" and len(wl) > 0
        print(f"[{'PASS' if tri_ok else 'FAIL'}] 三态判定（{st_clean}/{st_warn}/{st_fail}，告警 {len(wl)} 条）")
        ok &= tri_ok
        # docx 往返：article-format 出稿后标题是段落样式、正文无 `##`，
        # 这条断言防 parse_docx 退回「只认字面 ##」的旧缺陷。
        try:
            from docx import Document
        except ImportError:
            print("[SKIP] docx 往返（未安装 python-docx）")
        else:
            dx = Path(td) / "good.docx"
            doc = Document()

            def flush(rows):
                if not rows:
                    return
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                for i, r in enumerate(rows):
                    for j, c in enumerate(r[:len(rows[0])]):
                        t.cell(i, j).text = c

            doc.add_paragraph("会议纪要：测试会议", style="Title")
            pending = []
            for line in SELFTEST_OK.splitlines():
                s = line.strip()
                if s.startswith("# "):
                    continue
                is_table = s.startswith("|")
                if not is_table:
                    flush(pending)
                    pending = []
                if is_table:
                    cells = [c.strip() for c in s.strip("|").split("|")]
                    if not all(set(c) <= set("-: ") for c in cells):
                        pending.append(cells)
                    continue
                if s.startswith("### "):
                    doc.add_paragraph(s[4:], style="Heading 2")
                elif s.startswith("## "):
                    doc.add_paragraph(s[3:].strip(), style="Heading 1")
                elif s:
                    doc.add_paragraph(s)
            flush(pending)
            doc.save(str(dx))
            rc4, res4 = validate(dx)
            docx_ok = rc4 == 0
            print(f"[{'PASS' if docx_ok else 'FAIL'}] docx 往返按样式识别标题（exit={rc4}）")
            if not docx_ok:
                for level, name, detail in res4:
                    if level != "PASS":
                        print(f"  {level} {name} {detail}")
            ok &= docx_ok
    print("SELFTEST", "PASS" if ok else "FAIL")
    return 0 if ok else 1


def main(argv=None):
    ap = argparse.ArgumentParser(description="结构化总结结构校验与源文核对（会议版章节名）")
    ap.add_argument("file", nargs="?", help="纪要 md/docx")
    ap.add_argument("--source", default=None, help="源材料文本（转写稿/笔记）")
    ap.add_argument("--strict", action="store_true", help="WARN 也判失败")
    ap.add_argument("--json", dest="json_path", default=None, help="结果落 JSON")
    ap.add_argument("--selftest", action="store_true")
    ns = ap.parse_args(argv)
    if ns.selftest:
        return selftest()
    if not ns.file:
        ap.error("需要纪要文件，或 --selftest")
    path = Path(ns.file)
    if not path.is_file():
        print(f"文件不存在：{path}")
        return 2
    source = Path(ns.source).read_text(encoding="utf-8", errors="replace") if ns.source else None
    rc, results = validate(path, source, ns.strict)
    for level, name, detail in results:
        print(f"[{level}] {name}" + (f"  {detail}" if detail else ""))
    fails = sum(1 for r in results if r[0] == "FAIL")
    warns = sum(1 for r in results if r[0] == "WARN")
    status, warnings = summarize(results)
    print(f"\n结论：{'PASS' if rc == 0 else 'FAIL'}（FAIL {fails}，WARN {warns}）")
    print(f"状态：{status}")
    for w in warnings:
        print(f"  告警步骤：{w['step']} — {w['reason']}")
    if ns.json_path:
        Path(ns.json_path).write_text(json.dumps(
            {"file": str(path), "verdict": "PASS" if rc == 0 else "FAIL",
             "status": status,
             "results": [{"level": a, "name": b, "detail": c} for a, b, c in results],
             "processing_warnings": warnings},
            ensure_ascii=False, indent=2), encoding="utf-8")
    return rc


if __name__ == "__main__":
    raise SystemExit(main())
